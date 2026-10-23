import json,glob,collections
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
def load(c):
    r=[]
    for f in glob.glob(f'live/{c}/*.json'):
        d=json.load(open(f)); did=d.get('id'); d=d.get('data',d); d['id']=did or f.split('/')[-1][:-5]; r.append(d)
    return r
C,O,P=load('customers'),load('orders'),load('payments')
cur=lambda x:'TRY' if x.get('currency')=='TRY' else 'USD'
def tot(o):
    s=sum((float(i.get('qty') or 0))*(float(i.get('unitPrice') or 0)) for i in o.get('items',[]))
    return round(s*(1+float(o.get('vatRate') or 0)/100),2)
st=collections.defaultdict(lambda:{'n':0,'USD':[0,0],'TRY':[0,0],'lo':'','lp':''})
for o in O:
    if o.get('status')=='İptal': continue
    s=st[o['customerId']]; s[cur(o)][0]+=tot(o)
    if not o.get('kind'): s['n']+=1; s['lo']=max(s['lo'],o.get('date') or '')
for p in P:
    s=st[p['customerId']]; s[cur(p)][1]+=float(p.get('amount') or 0); s['lp']=max(s['lp'],p.get('date') or '')
hasTL=any(abs(s['TRY'][0])+abs(s['TRY'][1])>0.005 for s in st.values())
rows=[]
for c in C:
    s=st[c['id']]; rows.append((c,s))
rows.sort(key=lambda x:-(x[1]['USD'][0]-x[1]['USD'][1]))
wb=Workbook(); ws=wb.active; ws.title='Müşteriler'
F=lambda **k:Font(name='Arial',**k)
ws['A1']='Müşteri Listesi ve Bakiyeler'; ws['A1'].font=F(bold=True,size=14)
ws['A2']=f'Kaynak: Sipariş ve Tahsilat Defteri, 23.10.2026 itibarıyla. {len(C)} müşteri. Bakiye = Toplam satış − Tahsilat (iptal edilen siparişler hariç, devir kayıtları dahil).'
ws['A2'].font=F(italic=True,size=9,color='595959')
hdr=['Sıra','Müşteri / Firma','Yetkili','Telefon','E-posta','Şehir','Vergi D. / No','Adres','Sipariş sayısı','Son sipariş','Son tahsilat','Toplam satış (USD)','Tahsilat (USD)','Bakiye (USD)']
if hasTL: hdr+=['Toplam satış (TL)','Tahsilat (TL)','Bakiye (TL)']
hdr+=['Not']
H=4
for j,h in enumerate(hdr,1):
    c=ws.cell(H,j,h); c.font=F(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F3A5F'); c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
col={h:j for j,h in enumerate(hdr,1)}
L=lambda h:get_column_letter(col[h])
from datetime import date
def dt(s): 
    try: y,m,d=map(int,s.split('-')); return date(y,m,d)
    except: return None
thin=Side(style='thin',color='D9D9D9')
for i,(c,s) in enumerate(rows):
    r=H+1+i
    vals={'Sıra':i+1,'Müşteri / Firma':c.get('company') or c.get('name'),'Yetkili':c.get('name') if c.get('company') else '',
      'Telefon':c.get('phone',''),'E-posta':c.get('email',''),'Şehir':c.get('city',''),'Vergi D. / No':c.get('taxNo',''),'Adres':c.get('address',''),
      'Sipariş sayısı':s['n'],'Son sipariş':dt(s['lo']),'Son tahsilat':dt(s['lp']),
      'Toplam satış (USD)':round(s['USD'][0],2),'Tahsilat (USD)':round(s['USD'][1],2),
      'Bakiye (USD)':f"={L('Toplam satış (USD)')}{r}-{L('Tahsilat (USD)')}{r}",'Not':c.get('note','')}
    if hasTL: vals.update({'Toplam satış (TL)':round(s['TRY'][0],2),'Tahsilat (TL)':round(s['TRY'][1],2),'Bakiye (TL)':f"={L('Toplam satış (TL)')}{r}-{L('Tahsilat (TL)')}{r}"})
    for h,v in vals.items():
        cell=ws.cell(r,col[h],v); cell.font=F(); cell.border=Border(bottom=thin)
        if '(USD)' in h: cell.number_format='"$"#,##0.00;[Red]-"$"#,##0.00;"-"'
        elif '(TL)' in h: cell.number_format='#,##0.00" ₺";[Red]-#,##0.00" ₺";"-"'
        elif h in('Son sipariş','Son tahsilat'): cell.number_format='dd.mm.yyyy'
    ws.cell(r,col['Bakiye (USD)']).font=F(bold=True)
last=H+len(rows); T=last+1
ws.cell(T,col['Müşteri / Firma'],'TOPLAM').font=F(bold=True)
for h in hdr:
    if '(USD)' in h or '(TL)' in h or h=='Sipariş sayısı':
        c=ws.cell(T,col[h],f'=SUM({L(h)}{H+1}:{L(h)}{last})'); c.font=F(bold=True)
        c.number_format=ws.cell(H+1,col[h]).number_format
for j in range(1,len(hdr)+1): ws.cell(T,j).fill=PatternFill('solid',fgColor='E8EEF5'); ws.cell(T,j).border=Border(top=Side(style='thin',color='1F3A5F'))
W={'Sıra':6,'Müşteri / Firma':32,'Yetkili':16,'Telefon':14,'E-posta':20,'Şehir':12,'Vergi D. / No':14,'Adres':20,'Sipariş sayısı':10,'Son sipariş':12,'Son tahsilat':12,'Not':40}
for h,j in col.items(): ws.column_dimensions[get_column_letter(j)].width=W.get(h,15)
ws.row_dimensions[H].height=32
ws.freeze_panes=ws.cell(H+1,3); ws.auto_filter.ref=f'A{H}:{get_column_letter(len(hdr))}{last}'
wb.save('out/Musteri_Listesi.xlsx')
print(len(rows),'hasTL',hasTL, 'USD bakiye toplam', round(sum(s['USD'][0]-s['USD'][1] for _,s in rows),2))
print([(c.get('company'),round(s['USD'][0]-s['USD'][1],2)) for c,s in rows[:5]])
